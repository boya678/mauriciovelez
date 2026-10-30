import io
import json
import re
from datetime import date, datetime, timedelta, timezone
from typing import Optional
from zoneinfo import ZoneInfo

import openpyxl
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.admin_security import require_admin
from app.core.config import settings
from app.database import get_db
from app.models.audit_log import AuditLog
from app.models.rifa_enumerada import RifaEnumeradaAviso, RifaEnumeradaBase, RifaEnumeradaBoleta, RifaEnumeradaLote
from app.services.rifa_enumerada import (
    BASE_ID,
    MAX_MPS,
    hay_lote_enviando,
    iniciar_lote,
    notificar_retiro,
    renombrar_blob_retirada,
    retirar_boletas,
    verificar_blob_publico,
)
from app.services.rifa_enumerada_util import leer_celulares
from app.services.servicios_config import get_rifa_enumerada_config, set_rifa_enumerada_config

router = APIRouter(prefix="/admin/servicios/rifa-enumerada", tags=["Admin Servicios"])

COL_TZ = ZoneInfo("America/Bogota")
IMAGEN_MAX_BYTES = 5 * 1024 * 1024
EXCEL_MAX_BYTES = 5 * 1024 * 1024
ERRORES_VISIBLES = 200


class RifaEnumeradaOut(BaseModel):
    activo: bool
    valor: int
    cantidad: int
    tiene_base: bool


class RifaEnumeradaIn(BaseModel):
    activo: bool
    valor: int = Field(ge=0)
    cantidad: int = Field(ge=1, le=100)


class EnviarIn(BaseModel):
    incluir_repetidos: bool = False


class FilaInvalida(BaseModel):
    fila: int
    valor: str


class ErrorEnvio(BaseModel):
    celular: str
    error: Optional[str]


class LoteOut(BaseModel):
    id: int
    estado: str
    archivo: str
    total_filas: int
    validas: int
    invalidas: int
    duplicadas: int
    ya_recibieron: int
    invalidas_detalle: list[FilaInvalida]
    avisos_total: int
    avisos_enviados: int
    avisos_pendientes: int
    errores_total: int
    solicitaron: int
    boletas_enviadas: int
    errores: list[ErrorEnvio]


class BoletaOut(BaseModel):
    id: int
    numero: str
    estado_envio: str
    pagada: bool
    retirada: bool
    fecha: Optional[str]


class BoletasCelularOut(BaseModel):
    celular: str
    boletas: list[BoletaOut]


class RetirarIn(BaseModel):
    celular: str
    boleta_ids: list[int] = Field(min_length=1, max_length=100)


class RetirarOut(BaseModel):
    numeros: list[str]


def _celular_local(valor: str) -> str:
    digitos = re.sub(r"\D", "", valor)
    if len(digitos) == 12 and digitos.startswith("57"):
        digitos = digitos[2:]
    if len(digitos) != 10:
        raise HTTPException(status_code=400, detail="El celular debe tener 10 dígitos (con o sin 57)")
    return digitos


def _out(db: Session, cfg) -> RifaEnumeradaOut:
    return RifaEnumeradaOut(
        activo=cfg.activo,
        valor=cfg.valor,
        cantidad=cfg.cantidad,
        tiene_base=db.get(RifaEnumeradaBase, BASE_ID) is not None,
    )


def _celulares_atendidos(db: Session, celulares: list[str]) -> set[str]:
    """Celulares que ya recibieron un aviso o boletas en algún lote."""
    encontrados: set[str] = set()
    for i in range(0, len(celulares), 1000):
        trozo = celulares[i:i + 1000]
        encontrados.update(
            db.execute(
                select(RifaEnumeradaBoleta.celular).where(RifaEnumeradaBoleta.celular.in_(trozo)).distinct()
            ).scalars().all()
        )
        encontrados.update(
            db.execute(
                select(RifaEnumeradaAviso.celular)
                .where(RifaEnumeradaAviso.celular.in_(trozo), RifaEnumeradaAviso.estado == "enviado")
                .distinct()
            ).scalars().all()
        )
    return encontrados


def _lote_out(db: Session, lote: RifaEnumeradaLote) -> LoteOut:
    por_estado = dict(
        db.execute(
            select(RifaEnumeradaAviso.estado, func.count())
            .where(RifaEnumeradaAviso.lote_id == lote.id)
            .group_by(RifaEnumeradaAviso.estado)
        ).all()
    )
    errores = db.execute(
        select(RifaEnumeradaAviso.celular, RifaEnumeradaAviso.error)
        .where(RifaEnumeradaAviso.lote_id == lote.id, RifaEnumeradaAviso.estado == "error")
        .order_by(RifaEnumeradaAviso.id)
        .limit(ERRORES_VISIBLES)
    ).all()
    solicitaron = db.execute(
        select(func.count()).select_from(RifaEnumeradaAviso)
        .where(RifaEnumeradaAviso.lote_id == lote.id, RifaEnumeradaAviso.solicitada_at.is_not(None))
    ).scalar_one()
    boletas_enviadas = db.execute(
        select(func.count()).select_from(RifaEnumeradaBoleta)
        .where(RifaEnumeradaBoleta.lote_id == lote.id, RifaEnumeradaBoleta.estado_envio == "enviada")
    ).scalar_one()
    ya_recibieron = (
        len(_celulares_atendidos(db, lote.celulares.split())) if lote.estado == "cargado" else 0
    )
    return LoteOut(
        id=lote.id,
        estado=lote.estado,
        archivo=lote.archivo,
        total_filas=lote.total_filas,
        validas=lote.validas,
        invalidas=lote.invalidas,
        duplicadas=lote.duplicadas,
        ya_recibieron=ya_recibieron,
        invalidas_detalle=[FilaInvalida(**f) for f in json.loads(lote.invalidas_json or "[]")],
        avisos_total=sum(por_estado.values()),
        avisos_enviados=por_estado.get("enviado", 0),
        avisos_pendientes=por_estado.get("pendiente", 0),
        errores_total=por_estado.get("error", 0),
        solicitaron=solicitaron,
        boletas_enviadas=boletas_enviadas,
        errores=[ErrorEnvio(celular=c, error=e) for c, e in errores],
    )


@router.get("", response_model=RifaEnumeradaOut)
def get_rifa_enumerada(
    db: Session = Depends(get_db),
    _user=Depends(require_admin),
):
    return _out(db, get_rifa_enumerada_config(db))


@router.put("", response_model=RifaEnumeradaOut)
def put_rifa_enumerada(
    payload: RifaEnumeradaIn,
    db: Session = Depends(get_db),
    user=Depends(require_admin),
):
    if payload.activo and payload.valor <= 0:
        raise HTTPException(status_code=400, detail="Debes ingresar un monto mayor a 0 cuando el servicio esta activo")
    if payload.activo and db.get(RifaEnumeradaBase, BASE_ID) is None:
        raise HTTPException(status_code=400, detail="Debes subir la imagen base antes de activar el servicio")

    cfg = set_rifa_enumerada_config(db, payload.activo, payload.valor, payload.cantidad)
    db.add(AuditLog(
        platform_user_id=user.id,
        usuario=user.usuario,
        action="UPDATE",
        entity="servicios.rifa_enumerada",
        entity_id="rifa_enumerada",
        detail={"activo": cfg.activo, "valor": cfg.valor, "cantidad": cfg.cantidad},
    ))
    db.commit()
    return _out(db, cfg)


@router.post("/base", response_model=RifaEnumeradaOut)
def subir_imagen_base(
    imagen: UploadFile = File(...),
    db: Session = Depends(get_db),
    user=Depends(require_admin),
):
    from PIL import Image

    datos = imagen.file.read()
    es_jpg = datos.startswith(b"\xff\xd8\xff")
    es_png = datos.startswith(b"\x89PNG\r\n\x1a\n")
    if not (es_jpg or es_png):
        raise HTTPException(status_code=400, detail="La imagen debe ser JPG o PNG")
    if len(datos) > IMAGEN_MAX_BYTES:
        raise HTTPException(status_code=400, detail="La imagen no puede superar 5 MB")
    try:
        Image.open(io.BytesIO(datos)).verify()
    except Exception:
        raise HTTPException(status_code=400, detail="No se pudo leer la imagen")

    mime = "image/jpeg" if es_jpg else "image/png"
    existente = db.get(RifaEnumeradaBase, BASE_ID)
    if existente is None:
        db.add(RifaEnumeradaBase(id=BASE_ID, imagen_data=datos, imagen_mime=mime))
    else:
        existente.imagen_data = datos
        existente.imagen_mime = mime
        existente.updated_at = datetime.now(timezone.utc)
    db.add(AuditLog(
        platform_user_id=user.id,
        usuario=user.usuario,
        action="UPDATE",
        entity="servicios.rifa_enumerada",
        entity_id="rifa_enumerada",
        detail={"imagen_base_bytes": len(datos)},
    ))
    db.commit()
    return _out(db, get_rifa_enumerada_config(db))


@router.post("/lote", response_model=LoteOut)
def cargar_excel(
    archivo: UploadFile = File(...),
    db: Session = Depends(get_db),
    user=Depends(require_admin),
):
    datos = archivo.file.read()
    if len(datos) > EXCEL_MAX_BYTES:
        raise HTTPException(status_code=400, detail="El Excel no puede superar 5 MB")
    try:
        leido = leer_celulares(datos)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    except Exception:
        raise HTTPException(status_code=400, detail="No se pudo leer el Excel; debe ser un archivo .xlsx")
    if not leido["validas"]:
        raise HTTPException(status_code=400, detail="El Excel no tiene celulares v\u00e1lidos en la primera columna")

    lote = RifaEnumeradaLote(
        usuario=user.usuario,
        archivo=(archivo.filename or "excel.xlsx")[:200],
        celulares="\n".join(leido["validas"]),
        total_filas=leido["total_filas"],
        validas=len(leido["validas"]),
        invalidas=len(leido["invalidas"]),
        duplicadas=leido["duplicadas"],
        invalidas_json=json.dumps(leido["invalidas"][:20]),
    )
    db.add(lote)
    db.commit()
    db.refresh(lote)
    return _lote_out(db, lote)


@router.get("/lote/ultimo", response_model=Optional[LoteOut])
def ultimo_lote(
    db: Session = Depends(get_db),
    _user=Depends(require_admin),
):
    lote = db.execute(select(RifaEnumeradaLote).order_by(RifaEnumeradaLote.id.desc()).limit(1)).scalar_one_or_none()
    return _lote_out(db, lote) if lote else None


@router.get("/lote/{lote_id}", response_model=LoteOut)
def get_lote(
    lote_id: int,
    db: Session = Depends(get_db),
    _user=Depends(require_admin),
):
    lote = db.get(RifaEnumeradaLote, lote_id)
    if lote is None:
        raise HTTPException(status_code=404, detail="Lote no encontrado")
    return _lote_out(db, lote)


@router.post("/lote/{lote_id}/enviar", response_model=LoteOut)
def enviar_lote(
    lote_id: int,
    payload: EnviarIn,
    db: Session = Depends(get_db),
    user=Depends(require_admin),
):
    lote = db.get(RifaEnumeradaLote, lote_id)
    if lote is None:
        raise HTTPException(status_code=404, detail="Lote no encontrado")
    if lote.estado != "cargado":
        raise HTTPException(status_code=409, detail="Este lote ya fue enviado")
    if hay_lote_enviando(db):
        raise HTTPException(status_code=409, detail="Ya hay un env\u00edo en curso; espera a que termine")

    faltantes = [
        nombre for nombre, valor in (
            ("WHATSAPP_TEMPLATE_NOTIFICACION_RIFA_ENUMERADA", settings.WHATSAPP_TEMPLATE_NOTIFICACION_RIFA_ENUMERADA),
            ("RIFA_BLOB_STORAGE_CONNECTION_STRING", settings.RIFA_BLOB_STORAGE_CONNECTION_STRING),
            ("RIFA_CONTAINER", settings.RIFA_CONTAINER),
        ) if not valor
    ]
    if faltantes:
        raise HTTPException(status_code=503, detail=f"Falta configurar: {', '.join(faltantes)}")

    base = db.get(RifaEnumeradaBase, BASE_ID)
    if base is None:
        raise HTTPException(status_code=400, detail="Debes subir la imagen base")
    cfg = get_rifa_enumerada_config(db)

    celulares = lote.celulares.split()
    if not payload.incluir_repetidos:
        previos = _celulares_atendidos(db, celulares)
        celulares = [c for c in celulares if c not in previos]
    if not celulares:
        raise HTTPException(status_code=400, detail="Todos los celulares ya recibieron boletas")

    try:
        verificar_blob_publico(base)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Almacenamiento de im\u00e1genes: {exc}")

    ahora = datetime.now(timezone.utc)
    db.bulk_insert_mappings(RifaEnumeradaAviso, [
        {"lote_id": lote.id, "celular": c, "celular_local": c[-10:], "estado": "pendiente", "created_at": ahora}
        for c in celulares
    ])

    lote.estado = "enviando"
    lote.cantidad = cfg.cantidad
    lote.incluir_repetidos = payload.incluir_repetidos
    lote.iniciado_at = ahora
    db.add(AuditLog(
        platform_user_id=user.id,
        usuario=user.usuario,
        action="CREATE",
        entity="servicios.rifa_enumerada",
        entity_id=str(lote.id),
        detail={"avisos": len(celulares), "cantidad": cfg.cantidad, "max_mps": MAX_MPS},
    ))
    db.commit()

    iniciar_lote(lote.id)
    return _lote_out(db, lote)


@router.get("/celular/{celular}", response_model=BoletasCelularOut)
def boletas_de_celular(
    celular: str,
    db: Session = Depends(get_db),
    _user=Depends(require_admin),
):
    local = _celular_local(celular)
    boletas = db.execute(
        select(RifaEnumeradaBoleta).where(RifaEnumeradaBoleta.celular_local == local).order_by(RifaEnumeradaBoleta.id)
    ).scalars().all()
    if not boletas:
        raise HTTPException(status_code=404, detail="Ese celular no tiene boletas")
    return BoletasCelularOut(
        celular=f"57{local}",
        boletas=[
            BoletaOut(
                id=b.id,
                numero=b.numero,
                estado_envio=b.estado_envio,
                pagada=b.pagada,
                retirada=b.retirada,
                fecha=b.enviada_at.astimezone(COL_TZ).strftime("%d/%m/%Y %H:%M") if b.enviada_at else None,
            )
            for b in boletas
        ],
    )


@router.post("/retirar", response_model=RetirarOut)
def retirar(
    payload: RetirarIn,
    db: Session = Depends(get_db),
    user=Depends(require_admin),
):
    local = _celular_local(payload.celular)
    ids = sorted(set(payload.boleta_ids))
    boletas = retirar_boletas(db, local, ids, user.usuario)
    if len(boletas) != len(ids):
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail="Alguna boleta ya no se puede retirar (está pagada, ya fue retirada o no se envió). Consulta el celular de nuevo.",
        )
    numeros = [b.numero for b in boletas]
    db.add(AuditLog(
        platform_user_id=user.id,
        usuario=user.usuario,
        action="UPDATE",
        entity="servicios.rifa_enumerada",
        entity_id=local,
        detail={"retiro_boletas": numeros},
    ))
    db.commit()

    notificar_retiro(local, numeros)
    for b in boletas:
        renombrar_blob_retirada(b.celular, b.numero, b.id)
    return RetirarOut(numeros=numeros)


@router.get("/exportar")
def exportar_boletas(
    fecha_inicio: Optional[date] = Query(default=None),
    fecha_fin: Optional[date] = Query(default=None),
    _user=Depends(require_admin),
    db: Session = Depends(get_db),
):
    if fecha_inicio is None and fecha_fin is None:
        fecha_fin = datetime.now(COL_TZ).date()
        fecha_inicio = fecha_fin - timedelta(days=30)

    fecha_col = func.date(func.timezone("America/Bogota", RifaEnumeradaBoleta.created_at))
    q = select(RifaEnumeradaBoleta)
    if fecha_inicio:
        q = q.where(fecha_col >= fecha_inicio)
    if fecha_fin:
        q = q.where(fecha_col <= fecha_fin)
    boletas = db.execute(q.order_by(RifaEnumeradaBoleta.id)).scalars().all()

    def _fecha(valor):
        return valor.astimezone(COL_TZ).strftime("%d/%m/%Y %H:%M") if valor else ""

    estados = {"enviada": "Enviada", "pendiente": "Pendiente", "error": "No enviada"}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Boletas"
    ws.append(["Fecha env\u00edo", "Celular", "Boleta", "Estado env\u00edo", "Pagada", "Fecha pago", "Comprobante", "Monto pago", "Error", "Retirada", "Fecha retiro", "Retirada por"])
    for b in boletas:
        ws.append([
            _fecha(b.enviada_at or b.created_at),
            b.celular,
            b.numero,
            estados.get(b.estado_envio, b.estado_envio),
            "S\u00ed" if b.pagada else "No",
            _fecha(b.pagada_at),
            b.comprobante_num or "",
            float(b.pago_monto) if b.pago_monto is not None else "",
            b.error or "",
            "Sí" if b.retirada else "No",
            _fecha(b.retirada_at),
            b.retirada_por or "",
        ])
        ws.cell(row=ws.max_row, column=3).number_format = "@"

    ws2 = wb.create_sheet("Avisos")
    ws2.append(["Fecha aviso", "Celular", "Aviso", "Solicitó boletas", "Fecha solicitud", "Error"])
    avisos_q = select(RifaEnumeradaAviso)
    fecha_aviso_col = func.date(func.timezone("America/Bogota", RifaEnumeradaAviso.created_at))
    if fecha_inicio:
        avisos_q = avisos_q.where(fecha_aviso_col >= fecha_inicio)
    if fecha_fin:
        avisos_q = avisos_q.where(fecha_aviso_col <= fecha_fin)
    estados_aviso = {"enviado": "Enviado", "pendiente": "Pendiente", "error": "No enviado"}
    for a in db.execute(avisos_q.order_by(RifaEnumeradaAviso.id)).scalars().all():
        ws2.append([
            _fecha(a.enviado_at or a.created_at),
            a.celular,
            estados_aviso.get(a.estado, a.estado),
            "Sí" if a.solicitada_at else "No",
            _fecha(a.solicitada_at),
            a.error or "",
        ])

    salida = io.BytesIO()
    wb.save(salida)
    salida.seek(0)
    return StreamingResponse(
        salida,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=rifa_enumerada_{fecha_inicio}_a_{fecha_fin}.xlsx"},
    )
