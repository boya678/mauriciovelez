import io
from datetime import date, datetime, timedelta
from typing import Optional
from zoneinfo import ZoneInfo

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.core.admin_security import require_admin
from app.database import get_db
from app.models.audit_log import AuditLog
from app.models.cliente import Cliente
from app.models.loteria_resultado import LoteriaResultado
from app.models.numbers_historic import NumberHistoric
from app.models.numero_acierto import NumeroAcierto
from app.models.numero_rifa_entregado import NumeroRifaEntregado
from app.services.servicios_config import (
    get_conferencia_config,
    get_conferencia_vip_config,
    get_numero_relampago_config,
    get_relampago_personalizado_config,
    get_vip_config,
    set_conferencia_config,
    set_conferencia_vip_config,
    set_numero_relampago_config,
    set_relampago_personalizado_config,
    set_vip_config,
)

router = APIRouter(prefix="/admin/servicios", tags=["Admin Servicios"])


class NumeroRelampagoOut(BaseModel):
    activo: bool
    valor: int
    numero: str


class NumeroRelampagoIn(BaseModel):
    activo: bool
    valor: int = Field(ge=0)
    numero: str = ""


class VipOut(BaseModel):
    activo: bool
    valor: int


class VipIn(BaseModel):
    activo: bool
    valor: int = Field(ge=0)


class RelampagoPersonalizadoOut(BaseModel):
    activo: bool
    valor: int
    vigencia_dias: int
    rifa_activo: bool
    rifa_descripcion: str


class RelampagoPersonalizadoIn(BaseModel):
    activo: bool
    valor: int = Field(ge=0)
    vigencia_dias: int = Field(ge=1)
    rifa_activo: bool = False
    rifa_descripcion: str = Field(default="", max_length=200)


class ConferenciaOut(BaseModel):
    activo: bool
    valor: int
    fecha_aviso: str
    link_youtube: str


class ConferenciaIn(BaseModel):
    activo: bool
    valor: int = Field(ge=0)
    fecha_aviso: str = ""
    link_youtube: str = ""


class ConferenciaVipOut(BaseModel):
    activo: bool
    valor: int
    fecha_aviso: str
    link_youtube: str


class ConferenciaVipIn(BaseModel):
    activo: bool
    valor: int = Field(ge=0)
    fecha_aviso: str = ""
    link_youtube: str = ""


@router.get("/numero-relampago", response_model=NumeroRelampagoOut)
def get_numero_relampago(
    db: Session = Depends(get_db),
    _user=Depends(require_admin),
):
    cfg = get_numero_relampago_config(db)
    return NumeroRelampagoOut(activo=cfg.activo, valor=cfg.valor, numero=cfg.numero)


@router.put("/numero-relampago", response_model=NumeroRelampagoOut)
def put_numero_relampago(
    payload: NumeroRelampagoIn,
    db: Session = Depends(get_db),
    user=Depends(require_admin),
):
    numero = (payload.numero or "").strip()
    if payload.activo and not numero:
        raise HTTPException(status_code=400, detail="Debes ingresar el numero relampago cuando el servicio esta activo")

    cfg = set_numero_relampago_config(
        db=db,
        activo=payload.activo,
        valor=payload.valor,
        numero=numero,
    )
    db.add(AuditLog(
        platform_user_id=user.id,
        usuario=user.usuario,
        action="UPDATE",
        entity="servicios.numero_relampago",
        entity_id="numero_relampago",
        detail={"activo": cfg.activo, "valor": cfg.valor, "numero": cfg.numero},
    ))
    db.commit()

    return NumeroRelampagoOut(activo=cfg.activo, valor=cfg.valor, numero=cfg.numero)


@router.get("/vip", response_model=VipOut)
def get_vip(
    db: Session = Depends(get_db),
    _user=Depends(require_admin),
):
    cfg = get_vip_config(db)
    return VipOut(activo=cfg.activo, valor=cfg.valor)


@router.put("/vip", response_model=VipOut)
def put_vip(
    payload: VipIn,
    db: Session = Depends(get_db),
    user=Depends(require_admin),
):
    if payload.activo and payload.valor <= 0:
        raise HTTPException(status_code=400, detail="Debes ingresar un monto mayor a 0 cuando el servicio esta activo")

    cfg = set_vip_config(db=db, activo=payload.activo, valor=payload.valor)
    db.add(AuditLog(
        platform_user_id=user.id,
        usuario=user.usuario,
        action="UPDATE",
        entity="servicios.vip",
        entity_id="vip",
        detail={"activo": cfg.activo, "valor": cfg.valor},
    ))
    db.commit()

    return VipOut(activo=cfg.activo, valor=cfg.valor)


@router.get("/relampago-personalizado", response_model=RelampagoPersonalizadoOut)
def get_relampago_personalizado(
    db: Session = Depends(get_db),
    _user=Depends(require_admin),
):
    cfg = get_relampago_personalizado_config(db)
    return _relampago_personalizado_out(cfg)


def _relampago_personalizado_out(cfg) -> RelampagoPersonalizadoOut:
    return RelampagoPersonalizadoOut(
        activo=cfg.activo,
        valor=cfg.valor,
        vigencia_dias=cfg.vigencia_dias,
        rifa_activo=cfg.rifa_activo,
        rifa_descripcion=cfg.rifa_descripcion,
    )


@router.put("/relampago-personalizado", response_model=RelampagoPersonalizadoOut)
def put_relampago_personalizado(
    payload: RelampagoPersonalizadoIn,
    db: Session = Depends(get_db),
    user=Depends(require_admin),
):
    if payload.activo and payload.valor <= 0:
        raise HTTPException(status_code=400, detail="Debes ingresar un monto mayor a 0 cuando el servicio esta activo")
    descripcion = payload.rifa_descripcion.strip()
    if payload.activo and payload.rifa_activo and not descripcion:
        raise HTTPException(status_code=400, detail="Debes ingresar la descripcion del sorteo cuando el numero de rifa esta activo")

    cfg = set_relampago_personalizado_config(
        db=db,
        activo=payload.activo,
        valor=payload.valor,
        vigencia_dias=payload.vigencia_dias,
        rifa_activo=payload.rifa_activo,
        rifa_descripcion=descripcion,
    )
    db.add(AuditLog(
        platform_user_id=user.id,
        usuario=user.usuario,
        action="UPDATE",
        entity="servicios.relampago_personalizado",
        entity_id="relampago_personalizado",
        detail={
            "activo": cfg.activo,
            "valor": cfg.valor,
            "vigencia_dias": cfg.vigencia_dias,
            "rifa_activo": cfg.rifa_activo,
            "rifa_descripcion": cfg.rifa_descripcion,
        },
    ))
    db.commit()

    return _relampago_personalizado_out(cfg)


@router.get("/relampago-personalizado/rifa/exportar")
def exportar_numeros_rifa(
    fecha_inicio: Optional[date] = Query(default=None),
    fecha_fin: Optional[date] = Query(default=None),
    _user=Depends(require_admin),
    db: Session = Depends(get_db),
):
    col_tz = ZoneInfo("America/Bogota")
    if fecha_inicio is None and fecha_fin is None:
        fecha_fin = datetime.now(col_tz).date()
        fecha_inicio = fecha_fin - timedelta(days=30)

    fecha_col = func.date(func.timezone("America/Bogota", NumeroRifaEntregado.created_at))
    q = db.query(NumeroRifaEntregado, Cliente.nombre).outerjoin(
        Cliente, Cliente.id == NumeroRifaEntregado.cliente_id
    )
    if fecha_inicio:
        q = q.filter(fecha_col >= fecha_inicio)
    if fecha_fin:
        q = q.filter(fecha_col <= fecha_fin)
    filas = q.order_by(NumeroRifaEntregado.created_at.asc(), NumeroRifaEntregado.id.asc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Numeros de rifa"
    ws.append(["Fecha", "Número", "Celular", "Nombre", "Descripción del sorteo"])
    for registro, nombre in filas:
        ws.append([
            registro.created_at.astimezone(col_tz).strftime("%d/%m/%Y %H:%M"),
            registro.numero,
            registro.celular,
            nombre or "",
            registro.descripcion,
        ])
        ws.cell(row=ws.max_row, column=2).number_format = "@"

    # Hoja 2: números relampago asignados y sus aciertos (fecha y lotería hacia la derecha)
    asignados_q = db.query(NumberHistoric, Cliente.celular, Cliente.nombre).outerjoin(
        Cliente, Cliente.id == NumberHistoric.id_user
    ).filter(NumberHistoric.type == "relampago")
    if fecha_inicio:
        asignados_q = asignados_q.filter(NumberHistoric.date >= fecha_inicio)
    if fecha_fin:
        asignados_q = asignados_q.filter(NumberHistoric.date <= fecha_fin)
    asignados = asignados_q.order_by(NumberHistoric.date.asc(), NumberHistoric.id.asc()).all()

    aciertos_por_historic: dict[int, list[tuple[date, str]]] = {}
    if asignados:
        filas_aciertos = (
            db.query(NumeroAcierto.historic_id, LoteriaResultado.fecha, LoteriaResultado.loteria)
            .join(LoteriaResultado, LoteriaResultado.id == NumeroAcierto.resultado_id)
            .filter(NumeroAcierto.historic_id.in_([h.id for h, _, _ in asignados]))
            .order_by(LoteriaResultado.fecha.asc(), LoteriaResultado.loteria.asc())
            .all()
        )
        for historic_id, fecha_acierto, loteria in filas_aciertos:
            aciertos_por_historic.setdefault(historic_id, []).append((fecha_acierto, loteria))

    max_aciertos = max((len(v) for v in aciertos_por_historic.values()), default=0)
    encabezado = ["Fecha asignación", "Número", "Celular", "Nombre", "Aciertos"]
    for i in range(1, max_aciertos + 1):
        encabezado += [f"Acierto {i} - Fecha", f"Acierto {i} - Lotería"]

    ws2 = wb.create_sheet("Numeros asignados")
    ws2.append(encabezado)
    for historic, celular, nombre in asignados:
        aciertos = aciertos_por_historic.get(historic.id, [])
        fila = [
            historic.date.strftime("%d/%m/%Y"),
            historic.number,
            celular or "",
            nombre or "",
            len(aciertos),
        ]
        for fecha_acierto, loteria in aciertos:
            fila += [fecha_acierto.strftime("%d/%m/%Y"), loteria]
        ws2.append(fila)
        ws2.cell(row=ws2.max_row, column=2).number_format = "@"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"numeros_rifa_{fecha_inicio}_a_{fecha_fin}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/conferencia", response_model=ConferenciaOut)
def get_conferencia(
    db: Session = Depends(get_db),
    _user=Depends(require_admin),
):
    cfg = get_conferencia_config(db)
    return ConferenciaOut(
        activo=cfg.activo,
        valor=cfg.valor,
        fecha_aviso=cfg.fecha_aviso,
        link_youtube=cfg.link_youtube,
    )


@router.put("/conferencia", response_model=ConferenciaOut)
def put_conferencia(
    payload: ConferenciaIn,
    db: Session = Depends(get_db),
    user=Depends(require_admin),
):
    fecha_aviso = (payload.fecha_aviso or "").strip()
    link = (payload.link_youtube or "").strip()
    if payload.activo and not fecha_aviso:
        raise HTTPException(status_code=400, detail="Debes ingresar la fecha de aviso cuando conferencia esta activa")
    if payload.activo and not link:
        raise HTTPException(status_code=400, detail="Debes ingresar el link de YouTube cuando conferencia esta activa")

    cfg = set_conferencia_config(
        db=db,
        activo=payload.activo,
        valor=payload.valor,
        fecha_aviso=fecha_aviso,
        link_youtube=link,
    )
    db.add(AuditLog(
        platform_user_id=user.id,
        usuario=user.usuario,
        action="UPDATE",
        entity="servicios.conferencia",
        entity_id="conferencia",
        detail={
            "activo": cfg.activo,
            "valor": cfg.valor,
            "fecha_aviso": cfg.fecha_aviso,
            "link_youtube": cfg.link_youtube,
        },
    ))
    db.commit()

    return ConferenciaOut(
        activo=cfg.activo,
        valor=cfg.valor,
        fecha_aviso=cfg.fecha_aviso,
        link_youtube=cfg.link_youtube,
    )


@router.get("/conferencia-vip", response_model=ConferenciaVipOut)
def get_conferencia_vip(
    db: Session = Depends(get_db),
    _user=Depends(require_admin),
):
    cfg = get_conferencia_vip_config(db)
    return ConferenciaVipOut(
        activo=cfg.activo,
        valor=cfg.valor,
        fecha_aviso=cfg.fecha_aviso,
        link_youtube=cfg.link_youtube,
    )


@router.put("/conferencia-vip", response_model=ConferenciaVipOut)
def put_conferencia_vip(
    payload: ConferenciaVipIn,
    db: Session = Depends(get_db),
    user=Depends(require_admin),
):
    fecha_aviso = (payload.fecha_aviso or "").strip()
    link = (payload.link_youtube or "").strip()
    if payload.activo and not fecha_aviso:
        raise HTTPException(status_code=400, detail="Debes ingresar la fecha de aviso cuando conferencia VIP esta activa")
    if payload.activo and not link:
        raise HTTPException(status_code=400, detail="Debes ingresar el link de YouTube cuando conferencia VIP esta activa")

    cfg = set_conferencia_vip_config(
        db=db,
        activo=payload.activo,
        valor=payload.valor,
        fecha_aviso=fecha_aviso,
        link_youtube=link,
    )
    db.add(AuditLog(
        platform_user_id=user.id,
        usuario=user.usuario,
        action="UPDATE",
        entity="servicios.conferencia_vip",
        entity_id="conferencia_vip",
        detail={
            "activo": cfg.activo,
            "valor": cfg.valor,
            "fecha_aviso": cfg.fecha_aviso,
            "link_youtube": cfg.link_youtube,
        },
    ))
    db.commit()

    return ConferenciaVipOut(
        activo=cfg.activo,
        valor=cfg.valor,
        fecha_aviso=cfg.fecha_aviso,
        link_youtube=cfg.link_youtube,
    )
